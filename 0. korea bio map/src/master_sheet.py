"""Hand-entered data in one Excel file: data/about_bio_master.xlsx.

Tabs (people are named by 이름 + 소속, not IDs, so the sheet is easy to fill):

- 교수추가: professors / researchers missing from the crawl. Added to
  professors_seed.csv on the next run; ORCID or OpenAlex ID makes the
  identity match exact.
- 관계추가: lines to draw by hand. 관계 = 공동연구 / 지도교수-제자 /
  포닥멘토-포닥 (A is the PI). Hand-entered lines count as verified.
- 연구실정보: per-lab figures (graduation time, papers per student, lab
  size, recruiting) shown on the professor card. Aggregates only, never
  individual students.

Rows that cannot be matched to a professor are listed in
output/master_unmatched.csv.
"""
from __future__ import annotations

import re
from pathlib import Path

import pandas as pd

from config import DATA_DIR, OUTPUT_DIR

MASTER_PATH = DATA_DIR / "about_bio_master.xlsx"

PROFESSOR_SHEET = "교수추가"
RELATION_SHEET = "관계추가"
LAB_SHEET = "연구실정보"
GUIDE_SHEET = "안내"

PROFESSOR_HEADERS = ["이름", "영문이름", "소속", "학과", "ORCID", "OpenAlex ID", "홈페이지", "메모"]
RELATION_HEADERS = ["A 이름", "A 소속", "B 이름", "B 소속", "관계", "근거 링크", "메모"]
# (header, key on the web node, kind)
LAB_FIELDS = [
    ("이름", "", "key"),
    ("소속", "", "key"),
    ("박사 평균 졸업기간(년)", "phd_years", "num"),
    ("석박통합 평균 졸업기간(년)", "integrated_years", "num"),
    ("석사 평균 졸업기간(년)", "ms_years", "num"),
    ("박사 졸업생 수", "phd_graduates", "num"),
    ("학생당 평균 논문 수", "papers_per_student", "num"),
    ("학생당 평균 1저자 논문 수", "first_author_per_student", "num"),
    ("현재 박사과정", "phd_students", "num"),
    ("현재 석사과정", "ms_students", "num"),
    ("현재 포닥", "postdocs", "num"),
    ("모집 중", "recruiting", "text"),
    ("기준연도", "as_of", "text"),
    ("출처", "source", "text"),
    ("메모", "", "skip"),
]
LAB_HEADERS = [h for h, _, _ in LAB_FIELDS]
RELATION_TYPES = {
    "공동연구": "collaboration",
    "지도교수-제자": "advisor_student",
    "포닥멘토-포닥": "postdoc_mentor",
    # English values work too
    "collaboration": "collaboration",
    "advisor_student": "advisor_student",
    "postdoc_mentor": "postdoc_mentor",
}


def load_master(path: Path = MASTER_PATH) -> dict[str, pd.DataFrame]:
    if not path.exists():
        return {}
    sheets = pd.read_excel(path, sheet_name=None, dtype=str)
    return {name: df.fillna("").map(lambda v: str(v).strip()) for name, df in sheets.items()}


def _norm_org(value: str) -> str:
    return re.sub(r"[\s()·,.-]", "", str(value or "").split(";")[0]).casefold()


class ProfessorMatcher:
    """Find a professor_id from a name and an organisation written either in Korean or English."""

    def __init__(self, professors: pd.DataFrame, data_dir: Path = DATA_DIR):
        aliases: dict[str, str] = {}
        seed = data_dir / "universities_seed.csv"
        if seed.exists():
            unis = pd.read_csv(seed, dtype=str).fillna("")
            for _, r in unis.iterrows():
                canonical = _norm_org(r["university"])
                for v in (r["university"], r.get("name_ko", "")):
                    if v:
                        aliases[_norm_org(v)] = canonical
                        # 서울대 / 서울대학교 both work
                        aliases[_norm_org(str(v).replace("대학교", "대"))] = canonical
        self.aliases = aliases
        self.rows = [
            (str(r.get("professor_id", "")), str(r.get("name_ko", "")).strip(),
             str(r.get("name_en", "")).strip().casefold(), self._org(r.get("university", "")))
            for _, r in professors.iterrows()
        ]

    def _org(self, value: str) -> str:
        key = _norm_org(value)
        return self.aliases.get(key, key)

    def find(self, name: str, org: str) -> str:
        name = str(name or "").strip()
        if not name:
            return ""
        same_name = [r for r in self.rows if name in (r[1],) or name.casefold() == r[2]]
        if org:
            target = self._org(org)
            same_name = [r for r in same_name if r[3] == target or target in r[3] or r[3] in target]
        return same_name[0][0] if len(same_name) == 1 else ""


def add_master_professors(professors_path: Path, columns: list[str], path: Path = MASTER_PATH) -> int:
    """Append 교수추가 rows that are not in professors_seed.csv yet (same name + organisation)."""
    sheet = load_master(path).get(PROFESSOR_SHEET)
    if sheet is None or sheet.empty:
        return 0
    professors = pd.read_csv(professors_path, dtype=str).fillna("") if professors_path.exists() else pd.DataFrame(columns=columns)
    matcher = ProfessorMatcher(professors)
    nums = [int(m.group(1)) for p in professors.get("professor_id", []) if (m := re.fullmatch(r"P(\d+)", str(p)))]
    next_num = max(nums, default=0) + 1
    new = []
    for _, r in sheet.iterrows():
        name = r.get("이름", "") or r.get("영문이름", "")
        if not name or matcher.find(name, r.get("소속", "")):
            continue
        new.append({
            "professor_id": f"P{next_num:04d}",
            "name_ko": r.get("이름", ""),
            "name_en": r.get("영문이름", ""),
            "university": r.get("소속", ""),
            "department": r.get("학과", ""),
            "openalex_id": r.get("OpenAlex ID", ""),
            "source_url": r.get("홈페이지", ""),
            "orcid": r.get("ORCID", ""),
            "identity_status": "manual",
        })
        matcher.rows.append((new[-1]["professor_id"], new[-1]["name_ko"], new[-1]["name_en"].casefold(),
                             matcher._org(new[-1]["university"])))
        next_num += 1
    if new:
        out = pd.concat([professors, pd.DataFrame(new)], ignore_index=True)
        for c in columns:
            if c not in out.columns:
                out[c] = ""
        out[columns].fillna("").to_csv(professors_path, index=False, encoding="utf-8-sig")
    return len(new)


def master_relationships(professors: pd.DataFrame, path: Path = MASTER_PATH) -> pd.DataFrame:
    """관계추가 rows as relationship rows (verified, hand-entered)."""
    sheet = load_master(path).get(RELATION_SHEET)
    columns = ["relationship_id", "professor_a_id", "professor_b_id", "relationship_type",
               "evidence_url", "verified", "notes"]
    if sheet is None or sheet.empty:
        return pd.DataFrame(columns=columns)
    matcher = ProfessorMatcher(professors)
    rows, unmatched = [], []
    for i, r in sheet.iterrows():
        a = matcher.find(r.get("A 이름", ""), r.get("A 소속", ""))
        b = matcher.find(r.get("B 이름", ""), r.get("B 소속", ""))
        kind = RELATION_TYPES.get(r.get("관계", "").replace(" ", ""), "")
        if not (a and b and kind) or a == b:
            if any(r.get(h, "") for h in RELATION_HEADERS):
                unmatched.append({"sheet": RELATION_SHEET, "row": i + 2,
                                  "problem": "관계 값" if a and b else "교수를 찾지 못함 (이름+소속 확인)",
                                  "values": " / ".join(r.get(h, "") for h in RELATION_HEADERS[:5])})
            continue
        if kind == "collaboration":
            a, b = sorted([a, b])  # undirected, same order as automatic lines
        rows.append({"relationship_id": f"XLS{i + 2:04d}", "professor_a_id": a, "professor_b_id": b,
                     "relationship_type": kind, "evidence_url": r.get("근거 링크", ""),
                     "verified": "yes", "notes": r.get("메모", "") or "master sheet"})
    _report_unmatched(RELATION_SHEET, unmatched)
    return pd.DataFrame(rows, columns=columns)


def lab_info(professors: pd.DataFrame, path: Path = MASTER_PATH) -> dict[str, dict]:
    """{professor_id: {"phd_years": 5.5, "recruiting": "예", ...}} from 연구실정보."""
    sheet = load_master(path).get(LAB_SHEET)
    if sheet is None or sheet.empty:
        return {}
    matcher = ProfessorMatcher(professors)
    out, unmatched = {}, []
    for i, r in sheet.iterrows():
        pid = matcher.find(r.get("이름", ""), r.get("소속", ""))
        if not pid:
            if r.get("이름", ""):
                unmatched.append({"sheet": LAB_SHEET, "row": i + 2, "problem": "교수를 찾지 못함 (이름+소속 확인)",
                                  "values": f"{r.get('이름', '')} / {r.get('소속', '')}"})
            continue
        info = {}
        for header, key, kind in LAB_FIELDS:
            value = r.get(header, "")
            if not key or kind in ("key", "skip") or value == "":
                continue
            if kind == "num":
                try:
                    info[key] = round(float(value), 2)
                except ValueError:
                    continue
            else:
                info[key] = value
        if info:
            out[pid] = info
    _report_unmatched(LAB_SHEET, unmatched)
    return out


def _report_unmatched(sheet: str, rows: list[dict]) -> None:
    path = OUTPUT_DIR / "master_unmatched.csv"
    old = pd.read_csv(path, dtype=str).fillna("") if path.exists() else pd.DataFrame()
    if not old.empty:
        old = old[old["sheet"] != sheet]
    out = pd.concat([old, pd.DataFrame(rows, columns=["sheet", "row", "problem", "values"])], ignore_index=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out.to_csv(path, index=False, encoding="utf-8-sig")


def write_template(path: Path = MASTER_PATH) -> None:
    """Create the empty master workbook with headers, dropdowns and a guide tab."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    guide = wb.active
    guide.title = GUIDE_SHEET
    lines = [
        ("About Bio 마스터 시트", True),
        ("여기에 적은 내용은 매일 자동 실행(오전 10시) 때 지도에 반영됩니다.", False),
        ("사람은 ID 대신 '이름 + 소속'으로 적습니다. 소속은 한글(서울대학교, 서울대)이나 영문 모두 됩니다.", False),
        ("", False),
        ("[교수추가] 크롤링에서 빠진 교수·연구원. ORCID나 OpenAlex ID를 알면 적어주세요 (정확도↑).", False),
        ("[관계추가] 직접 긋는 선. 관계: 공동연구 / 지도교수-제자 / 포닥멘토-포닥 (A가 지도교수·멘토).", False),
        ("[연구실정보] 교수 카드에 표시되는 연구실 통계. 학생 개인 정보는 적지 말고 평균·인원수만 적습니다.", False),
        ("", False),
        ("예) 관계추가: A 이름=김철수, A 소속=서울대학교, B 이름=이영희, B 소속=KAIST, 관계=지도교수-제자", False),
        ("예) 연구실정보: 이름=김철수, 소속=서울대, 박사 평균 졸업기간(년)=5.5, 학생당 평균 논문 수=4, 모집 중=예", False),
        ("", False),
        ("찾지 못한 행은 output/master_unmatched.csv에 이유와 함께 남습니다.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        guide.cell(row=i, column=1, value=text).font = Font(bold=bold, size=14 if bold else 11)
    guide.column_dimensions["A"].width = 110

    head_fill = PatternFill("solid", fgColor="1F2A44")
    head_font = Font(bold=True, color="FFFFFF")
    for title, headers in ((PROFESSOR_SHEET, PROFESSOR_HEADERS), (RELATION_SHEET, RELATION_HEADERS),
                           (LAB_SHEET, LAB_HEADERS)):
        ws = wb.create_sheet(title)
        for col, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=col, value=h)
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[c.column_letter].width = max(12, len(h) * 2 + 2)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
        if title == RELATION_SHEET:
            dv = DataValidation(type="list", formula1='"공동연구,지도교수-제자,포닥멘토-포닥"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add("E2:E2000")
        if title == LAB_SHEET:
            dv = DataValidation(type="list", formula1='"예,아니오"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{ws.cell(row=1, column=LAB_HEADERS.index('모집 중') + 1).column_letter}2:"
                   f"{ws.cell(row=1, column=LAB_HEADERS.index('모집 중') + 1).column_letter}2000")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    write_template()
    print(f"Wrote {MASTER_PATH}")
